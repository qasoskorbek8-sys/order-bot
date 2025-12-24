from openpyxl import Workbook, load_workbook
import os
import threading
import time

FILE_NAME = "data/orders.xlsx"
lock = threading.Lock()

def save_to_excel(data: dict):
    with lock:
        for _ in range(3):  # 3 marta urinadi
            try:
                if not os.path.exists(FILE_NAME):
                    wb = Workbook()
                    ws = wb.active
                    ws.append(list(data.keys()))
                else:
                    wb = load_workbook(FILE_NAME)
                    ws = wb.active

                ws.append(list(data.values()))
                wb.save(FILE_NAME)
                wb.close()
                return
            except PermissionError:
                time.sleep(0.5)

def update_status(order_id, status):
    with lock:
        for _ in range(3):
            try:
                wb = load_workbook(FILE_NAME)
                ws = wb.active

                for row in ws.iter_rows(min_row=2):
                    if row[0].value == order_id:
                        row[-1].value = status
                        break

                wb.save(FILE_NAME)
                wb.close()
                return
            except PermissionError:
                time.sleep(0.5)
